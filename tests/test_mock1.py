"""Builds model answers for Mock 1 and checks that they get full marks, and that an empty folder gets zero."""
import copy
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
from lxml import etree
from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PI

from helpers import add_field, fake_pdf, para, recalc_xlsx, set_margins, shade
from skilltest.content import mock1 as M
from skilltest.content.starters import write_pictures
from skilltest.core import grade_tasks
from skilltest.session import prepare_folder

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def solve_w1(f):
    d = Document()
    para(d, M.N_HEAD, bold=True, underline=True, size=16, align="CENTER")
    para(d, M.N_P1, align="JUSTIFY", spacing=1.5)
    para(d, M.N_P2, align="JUSTIFY", spacing=1.5)
    for i in M.N_ITEMS:
        d.add_paragraph(i, style="List Bullet")
    para(d, M.N_SIGN, italic=True, align="RIGHT")
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    set_margins(s, 0.5)
    fp = s.footer.paragraphs[0]
    fp.text = "Administration Division, ISRO"
    add_field(s.footer.add_paragraph(), "PAGE", "1")
    d.save(f / "Notice.docx")
    fake_pdf(f / "Notice.pdf")


def solve_w2(f):
    d = Document()
    para(d, M.S_TITLE, bold=True, size=14, font="Arial")
    t = d.add_table(rows=1, cols=4)
    t.style = d.styles["Table Grid"]
    for i, h in enumerate(["S.No", "Name", "Designation", "Basic Pay (₹)"]):
        c = t.rows[0].cells[i]
        c.paragraphs[0].add_run(h).bold = True
        shade(c)
    for r in sorted(M.STAFF, key=lambda r: r[1]):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    last = t.add_row().cells
    m = last[0].merge(last[2])
    m.text = "Total"
    add_field(last[3].paragraphs[0], "=SUM(ABOVE)", "143500")
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    s.header.paragraphs[0].text = "Confidential"
    d.save(f / "Staff.docx")


def solve_e1(f):
    wb = load_workbook(f / "Marks.xlsx")
    ws = wb.active
    ws.insert_rows(1)
    ws["A1"] = "CLASS RESULT 2026"
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    for c, h in zip("EFG", ["Total", "Average", "Result"]):
        ws[f"{c}2"] = h
    rows = sorted([[ws.cell(r, c).value for c in range(1, 5)] for r in range(3, 13)],
                  key=lambda r: -(r[1] + r[2] + r[3]))
    thin = Side(style="thin")
    for i, r in enumerate(rows):
        n = 3 + i
        for c, v in enumerate(r, 1):
            ws.cell(n, c, v)
        ws[f"E{n}"] = f"=SUM(B{n}:D{n})"
        ws[f"F{n}"] = f"=AVERAGE(B{n}:D{n})"
        ws[f"F{n}"].number_format = "0.00"
        ws[f"G{n}"] = f'=IF(F{n}>=40,"Pass","Fail")'
    for row in ws["A2:G12"]:
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[2]:
        if c.column <= 7:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFFF00")
    ws.freeze_panes = "A3"
    ch = BarChart()
    ch.type = "col"
    ch.title = "Total Marks"
    ch.add_data(Reference(ws, min_col=5, min_row=2, max_row=12), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=12))
    ws.add_chart(ch, "I2")
    wb.save(f / "Result.xlsx")
    recalc_xlsx(f / "Result.xlsx")


def solve_e2(f):
    wb = load_workbook(f / "Sales.xlsx")
    ws = wb["Data"]
    ws["F1"] = "Amount"
    for r in range(2, 22):
        ws[f"F{r}"] = f"=D{r}*E{r}"
        ws[f"F{r}"].number_format = '"₹"#,##0.00'
    ws.column_dimensions["C"].width = 20.7
    ws.conditional_formatting.add("F2:F21", CellIsRule(operator="greaterThan", formula=["50000"],
                                                        fill=PatternFill("solid", fgColor="FFC7CE")))
    ws["G1"] = "Status"
    dv = DataValidation(type="list", formula1='"Paid,Pending"')
    dv.add("G2:G21")
    ws.add_data_validation(dv)
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:1"
    ws.oddHeader.center.text = "Sales Report"
    s = wb.create_sheet("Summary")
    s["A1"], s["B1"] = "Region", "Total Sales"
    for i, rg in enumerate(M.REGIONS, 2):
        s[f"A{i}"] = rg
        s[f"B{i}"] = f"=SUMIF(Data!A:A,A{i},Data!F:F)"
    s.sheet_properties.tabColor = "FF0000"
    pc = PieChart()
    pc.title = "Region-wise Sales"
    pc.add_data(Reference(s, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    pc.set_categories(Reference(s, min_col=1, min_row=2, max_row=5))
    s.add_chart(pc, "D2")
    wb.save(f / "Sales.xlsx")
    recalc_xlsx(f / "Sales.xlsx")


def _set_theme_name(prs, name="Facet"):
    for m in prs.slide_masters:
        for rel in m.part.rels.values():
            if rel.reltype.endswith("/theme"):
                part = rel.target_part
                x = part.blob.decode().replace('name="Office Theme"', f'name="{name}"')
                part._blob = x.encode()


def _transition(slide):
    el = etree.SubElement(slide._element, f"{{{P_NS}}}transition")
    etree.SubElement(el, f"{{{P_NS}}}fade")


def _footer_ph(slide, text):
    for ph in slide.slide_layout.placeholders:
        t = str(ph.placeholder_format.type)
        if "FOOTER" in t or "SLIDE_NUMBER" in t:
            el = copy.deepcopy(ph._element)
            slide.shapes._spTree.append(el)
    for sh in slide.placeholders:
        if "FOOTER" in str(sh.placeholder_format.type):
            sh.text_frame.text = text


def solve_p1(f):
    prs = Presentation()
    _set_theme_name(prs)
    L = prs.slide_layouts
    s = prs.slides.add_slide(L[0]); s.shapes.title.text = "Indian Space Research Organisation"; s.placeholders[1].text = "Major Missions"
    s = prs.slides.add_slide(L[1]); s.shapes.title.text = "Recent Missions"; s.placeholders[1].text = "Chandrayaan-3\nAditya-L1\nGaganyaan"
    s = prs.slides.add_slide(L[3]); s.shapes.title.text = "Launch Vehicles"; s.placeholders[1].text = "PSLV and GSLV"
    s.shapes.add_picture(str(f / "Pictures" / "Rocket.png"), PI(5), PI(2))
    s = prs.slides.add_slide(L[5]); s.shapes.title.text = "Mission Table"
    t = s.shapes.add_table(3, 3, PI(1), PI(2), PI(6), PI(2)).table
    for r, row in enumerate([["Mission", "Year", "Purpose"], ["Chandrayaan-3", "2023", "Moon"], ["Aditya-L1", "2023", "Sun"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    s = prs.slides.add_slide(L[5]); s.shapes.title.text = "Thank You"
    for i, s in enumerate(prs.slides):
        _transition(s)
        if i:
            _footer_ph(s, "ISRO Skill Test")
    prs.save(f / "ISRO.pptx")


def solve_p2(f):
    prs = Presentation()
    L = prs.slide_layouts
    s = prs.slides.add_slide(L[0]); s.shapes.title.text = "Office Automation Training"; s.placeholders[1].text = "Skill Test Preparation"
    s.notes_slide.notes_text_frame.text = "Welcome all participants"
    s2 = prs.slides.add_slide(L[5]); s2.shapes.title.text = "Training Process"
    gf = etree.fromstring(
        f'<p:graphicFrame xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:nvGraphicFramePr><p:cNvPr id="9" name="Diagram"/>'
        '<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></p:xfrm>'
        '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"/></a:graphic></p:graphicFrame>')
    s2.shapes._spTree.append(gf)
    s3 = prs.slides.add_slide(L[6])
    g = s3.shapes.add_group_shape()
    r = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, PI(1), PI(1), PI(2), PI(1))
    r.fill.gradient()
    g.shapes.add_shape(MSO_SHAPE.OVAL, PI(4), PI(1), PI(2), PI(1.5))
    tb = s3.shapes.add_textbox(PI(1), PI(4), PI(3), PI(1)); tb.text_frame.text = "Thank You"
    s3._element.append(etree.fromstring(
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" presetClass="entr"/></p:par></p:tnLst></p:timing>'))
    for s in prs.slides:
        cSld = s._element.find(f"{{{P_NS}}}cSld")
        bg = etree.fromstring(f'<p:bg xmlns:p="{P_NS}" xmlns:a="{A_NS}"><p:bgPr><a:gradFill><a:gsLst>'
                              '<a:gs pos="0"><a:srgbClr val="FFFFFF"/></a:gs><a:gs pos="100000"><a:srgbClr val="4472C4"/></a:gs>'
                              '</a:gsLst></a:gradFill><a:effectLst/></p:bgPr></p:bg>')
        cSld.insert(0, bg)
    prs.save(f / "Training.pptx")
    fake_pdf(f / "Training.pdf")


def run():
    tmp = Path(tempfile.mkdtemp())
    try:
        tasks = M.PAPER.tasks
        empty = tmp / "empty"
        prepare_folder(empty, tasks)
        res = grade_tasks(tasks, empty)
        print("Empty/starter-only folder:", sum(r.got for r in res))
        for r in res:
            for c in r.checks:
                if c.ok:
                    print("   unexpected pass:", r.id, c.desc, "|", c.found)

        full = tmp / "full"
        prepare_folder(full, tasks)
        for s in (solve_w1, solve_w2, solve_e1, solve_e2, solve_p1, solve_p2):
            s(full)
        res = grade_tasks(tasks, full)
        bad = 0
        for r in res:
            print(f"{r.id}: {r.got}/{r.marks}")
            for c in r.checks:
                if not c.ok:
                    bad += 1
                    print("   FAIL:", c.desc, "|", c.found)
        print("TOTAL", sum(r.got for r in res), "/ 100")
        return bad
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(1 if run() else 0)
