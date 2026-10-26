"""Solve the formula-heavy Excel tasks of Mocks 2-5 with real formulas (recalculated by LibreOffice)
and check that they get full marks."""
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from helpers import recalc_xlsx
from skilltest.content import mock2, mock3, mock4, mock5
from skilltest.core import grade_task, all_file_names
from skilltest.session import prepare_folder


def m2e1(f):
    wb = load_workbook(f / "Salary.xlsx"); ws = wb["Pay"]
    ws["H1"], ws["I1"], ws["H2"], ws["I2"] = "DA Rate", 0.46, "HRA Rate", 0.24
    ws["I1"].number_format = ws["I2"].number_format = "0%"
    for r in range(2, 10):
        ws[f"D{r}"] = f"=C{r}*$I$1"; ws[f"E{r}"] = f"=C{r}*$I$2"; ws[f"F{r}"] = f"=ROUND(C{r}+D{r}+E{r},0)"
        for c in "CDEF":
            ws[f"{c}{r}"].number_format = "#,##0"
    ws["F11"] = "=MAX(F2:F9)"; ws["F12"] = "=MIN(F2:F9)"; ws["F13"] = "=COUNTIF(C2:C9,25500)"
    ws["I5"] = "E104"; ws["I6"] = "=VLOOKUP(I5,A2:F9,2,FALSE)"
    for c in ws[1][:6]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 30
    wb.save(f / "Salary.xlsx"); recalc_xlsx(f / "Salary.xlsx")


def m3e1(f):
    wb = load_workbook(f / "Attendance.xlsx"); ws = wb["March"]; ws.title = "March 2026"
    for i in range(20):
        ws.cell(1, 2 + i, i + 1)
        ws.column_dimensions[ws.cell(1, 2 + i).column_letter].width = 4.7
    ws["V1"], ws["W1"] = "Present", "Percent"
    for r in range(2, 12):
        ws[f"V{r}"] = f'=COUNTIF(B{r}:U{r},"P")'; ws[f"W{r}"] = f"=V{r}/20"; ws[f"W{r}"].number_format = "0%"
    for c in ws[1][:23]:
        c.font = Font(bold=True)
        if c.column > 1:
            c.alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("W2:W11", CellIsRule(operator="lessThan", formula=["0.75"], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws["Y1"] = "=TODAY()"; ws.freeze_panes = "B2"
    ch = BarChart(); ch.type = "bar"; ch.title = "Attendance %"
    ch.add_data(Reference(ws, min_col=23, min_row=1, max_row=11), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=11))
    from openpyxl.chart.label import DataLabelList
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ws.add_chart(ch, "A14")
    wb.save(f / "Attendance.xlsx"); recalc_xlsx(f / "Attendance.xlsx")


def m4e1(f):
    wb = load_workbook(f / "Bill.xlsx"); ws = wb["Bills"]
    rows = sorted([[ws.cell(r, c).value for c in (1, 2, 3)] for r in range(2, 10)])
    ws["D1"], ws["E1"], ws["F1"] = "Units", "Charge", "Total"
    thin = Side(style="thin")
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
        ws[f"D{i}"] = f"=C{i}-B{i}"
        ws[f"E{i}"] = f"=IF(D{i}<=100,D{i}*3,IF(D{i}<=200,300+(D{i}-100)*5,800+(D{i}-200)*7))"
        ws[f"F{i}"] = f"=ROUND(E{i}*1.05,0)"
        ws[f"E{i}"].number_format = ws[f"F{i}"].number_format = '"₹"#,##0.00'
    ws["F11"] = "=SUM(F2:F9)"; ws["D12"] = "=AVERAGE(D2:D9)"; ws["D13"] = "=MAX(D2:D9)"
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0"); dv.add("C2:C9"); ws.add_data_validation(dv)
    for row in ws["A1:F9"]:
        for c in row:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1][:6]:
        c.font = Font(bold=True)
    wb.save(f / "Bill.xlsx"); recalc_xlsx(f / "Bill.xlsx")


def m5e1(f):
    wb = load_workbook(f / "Exam.xlsx"); ws = wb["Sheet1"]; ws.title = "Result"
    ws["E1"], ws["F1"], ws["G1"] = "Total", "Percent", "Grade"
    for r in range(2, 12):
        ws[f"E{r}"] = f"=SUM(C{r}:D{r})"; ws[f"F{r}"] = f"=E{r}/200"; ws[f"F{r}"].number_format = "0.00%"
        ws[f"G{r}"] = f'=IF(F{r}>=0.75,"A",IF(F{r}>=0.6,"B",IF(F{r}>=0.4,"C","F")))'
    ws["I1"], ws["J1"], ws["I2"], ws["J2"] = "Number of A", '=COUNTIF(G2:G11,"A")', "Class Average", "=AVERAGE(E2:E11)"
    ws.conditional_formatting.add("G2:G11", CellIsRule(operator="equal", formula=['"F"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.freeze_panes = "A2"
    ch = LineChart(); ch.title = "Paper-wise Marks"
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=1, max_row=11), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=2, max_row=11)); ws.add_chart(ch, "L2")
    ws.auto_filter.ref = "A1:G11"; ws.auto_filter.add_filter_column(6, ["A"])
    wb.save(f / "Exam.xlsx"); recalc_xlsx(f / "Exam.xlsx")


def m5e2(f):
    wb = load_workbook(f / "Payroll.xlsx"); del wb["Notes"]; ws = wb["Jan"]
    ws["D1"], ws["E1"] = "Wage", "Per Day"
    for r in range(2, 8):
        ws[f"D{r}"] = f"=B{r}*C{r}"; ws[f"E{r}"] = f"=IFERROR(D{r}/B{r},0)"
    ws.column_dimensions["C"].hidden = True
    ws.oddHeader.center.text = "Payroll January"; ws.print_title_rows = "1:1"; ws.page_setup.orientation = "landscape"
    cp = wb.copy_worksheet(ws); cp.title = "Feb"
    ws.protection.sheet = True
    wb.save(f / "Payroll.xlsx"); recalc_xlsx(f / "Payroll.xlsx")
    import shutil as sh
    sh.copy(f / "Payroll.xlsx", f / "Payroll_tmp.xlsx")
    # fake an .xltx by changing the content type
    import zipfile
    zin = zipfile.ZipFile(f / "Payroll_tmp.xlsx")
    zout = zipfile.ZipFile(f / "Payroll.xltx", "w")
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "[Content_Types].xml":
            data = data.replace(b"spreadsheetml.sheet.main", b"spreadsheetml.template.main")
        zout.writestr(item, data)
    zout.close(); zin.close(); (f / "Payroll_tmp.xlsx").unlink()


def run():
    tmp = Path(tempfile.mkdtemp())
    bad = 0
    try:
        for mod, tid, solve in ((mock2, "M2-E1", m2e1), (mock3, "M3-E1", m3e1), (mock4, "M4-E1", m4e1),
                                (mock5, "M5-E1", m5e1), (mock5, "M5-E2", m5e2)):
            task = next(t for t in mod.PAPER.tasks if t.id == tid)
            f = tmp / tid
            prepare_folder(f, [task])
            solve(f)
            r = grade_task(task, f, all_file_names(mod.PAPER.tasks))
            print(f"{tid}: {r.got}/{r.marks}")
            for c in r.checks:
                if not c.ok:
                    bad += 1
                    print("   FAIL:", c.desc, "|", c.found)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return bad


if __name__ == "__main__":
    sys.exit(1 if run() else 0)
