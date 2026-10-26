"""Reading a .docx the way Word shows it, and the MS Word check builders."""
from __future__ import annotations

import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from .core import Check, norm, short, similarity

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSMAP = {"w": W_NS,
         "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
         "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
         "v": "urn:schemas-microsoft-com:vml",
         "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"}

TRUE_VALS = (None, "1", "true", "on")
ALIGN_NAME = {"left": "Left", "start": "Left", "center": "Center", "right": "Right", "end": "Right",
              "both": "Justify", "distribute": "Justify"}
COLOR_NAMES = {
    "red": ("FF0000", "C00000", "E00000"), "blue": ("0000FF", "0070C0", "002060", "4472C4", "1F4E79", "2E74B5"),
    "green": ("00B050", "00FF00", "008000", "70AD47", "548235"), "yellow": ("FFFF00",),
    "orange": ("FFC000", "ED7D31", "FF6600", "C55A11"), "purple": ("7030A0",),
}


def color_matches(hexval, want):
    """want = 'any' or a colour name or a hex string."""
    if not hexval or hexval.lower() in ("auto", "000000"):
        return False if want != "black" else True
    if want == "any":
        return True
    hexval = hexval.upper()
    if want in COLOR_NAMES:
        if hexval in COLOR_NAMES[want]:
            return True
        r, g, b = int(hexval[0:2], 16), int(hexval[2:4], 16), int(hexval[4:6], 16)
        dom = {"red": r > 150 and g < 90 and b < 90, "blue": b > 120 and r < 90 and b > g,
               "green": g > 100 and r < 120 and b < 120, "yellow": r > 200 and g > 200 and b < 100,
               "orange": r > 200 and 80 < g < 200 and b < 80, "purple": r > 80 and b > 100 and g < 80}
        return dom.get(want, False)
    return hexval == want.upper()


class WordDoc:
    def __init__(self, path):
        self.path = path
        self.doc = Document(str(path))
        self.body = self.doc.element.body
        self._styles = {}
        st = self.doc.styles.element
        for s in st.findall(qn("w:style")):
            self._styles[s.get(qn("w:styleId"))] = s
        self._default_para = None
        for s in st.findall(qn("w:style")):
            if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) in ("1", "true"):
                self._default_para = s.get(qn("w:styleId"))
        dd = st.find(qn("w:docDefaults"))
        self._rpr_default = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}") if dd is not None else None
        self._ppr_default = dd.find(f"{qn('w:pPrDefault')}/{qn('w:pPr')}") if dd is not None else None
        self._theme_fonts = self._read_theme_fonts()
        self._zip = zipfile.ZipFile(str(path))

    # ---------- raw parts
    def part_xml(self, name):
        try:
            return self._zip.read(name).decode("utf8", "ignore")
        except KeyError:
            return ""

    def parts(self, prefix):
        return [n for n in self._zip.namelist() if n.startswith(prefix)]

    def _read_theme_fonts(self):
        out = {}
        try:
            for rel in self.doc.part.rels.values():
                if rel.reltype.endswith("/theme"):
                    root = etree.fromstring(rel.target_part.blob)
                    a = "{%s}" % NSMAP["a"]
                    for kind in ("major", "minor"):
                        el = root.find(f".//{a}{kind}Font/{a}latin")
                        if el is not None:
                            out[kind] = el.get("typeface")
        except Exception:
            pass
        return out

    # ---------- paragraphs and text
    def paragraphs(self):
        return [Paragraph(p, self.doc._body) for p in self.body.iter(qn("w:p"))]

    def text(self):
        return "\n".join(p_text(p) for p in self.paragraphs())

    def find_para(self, text, fuzzy=True):
        want = norm(text)
        paras = self.paragraphs()
        for p in paras:
            if want and want in norm(p_text(p)):
                return p
        if fuzzy and len(want) > 12:
            best, score = None, 0.0
            for p in paras:
                t = norm(p_text(p))
                if not t:
                    continue
                s = similarity(want, t[: len(want) + 40])
                if s > score:
                    best, score = p, s
            if score >= 0.8:
                return best
        return None

    def runs_for(self, para, sub=None):
        runs, spans, pos = [], [], 0
        for r in para._p.iter(qn("w:r")):
            run = Run(r, para)
            t = run.text
            runs.append(run)
            spans.append((pos, pos + len(t)))
            pos += len(t)
        full = "".join(r.text for r in runs)
        if sub:
            i = norm_find(full, sub)
            if i is not None:
                a, b = i
                return [r for r, (s, e) in zip(runs, spans) if s < b and e > a and r.text.strip()]
        return [r for r in runs if r.text.strip()]

    # ---------- style chains
    def _style_chain(self, sid):
        out, seen = [], set()
        while sid and sid not in seen and sid in self._styles:
            seen.add(sid)
            s = self._styles[sid]
            out.append(s)
            b = s.find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None
        return out

    def para_style_id(self, para):
        ppr = para._p.pPr
        if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
            return ppr.find(qn("w:pStyle")).get(qn("w:val"))
        return self._default_para

    def style_name(self, para):
        sid = self.para_style_id(para)
        s = self._styles.get(sid)
        if s is not None and s.find(qn("w:name")) is not None:
            return s.find(qn("w:name")).get(qn("w:val"))
        return sid or "Normal"

    def rpr_chain(self, run, para):
        chain = []
        rpr = run._r.rPr
        if rpr is not None:
            chain.append(rpr)
            rs = rpr.find(qn("w:rStyle"))
            if rs is not None:
                chain += [s.find(qn("w:rPr")) for s in self._style_chain(rs.get(qn("w:val")))]
        chain += [s.find(qn("w:rPr")) for s in self._style_chain(self.para_style_id(para))]
        chain.append(self._rpr_default)
        return [c for c in chain if c is not None]

    def ppr_chain(self, para):
        chain = [para._p.pPr]
        chain += [s.find(qn("w:pPr")) for s in self._style_chain(self.para_style_id(para))]
        chain.append(self._ppr_default)
        return [c for c in chain if c is not None]

    def rprop(self, run, para, tag):
        for rpr in self.rpr_chain(run, para):
            el = rpr.find(qn(tag))
            if el is not None:
                return el
        return None

    def pprop(self, para, tag, attr=None):
        for ppr in self.ppr_chain(para):
            el = ppr.find(qn(tag))
            if el is not None:
                if attr is None:
                    return el
                v = el.get(qn(attr))
                if v is not None:
                    return v
        return None

    # ---------- run properties
    def is_on(self, run, para, tag):
        el = self.rprop(run, para, tag)
        return el is not None and el.get(qn("w:val")) in TRUE_VALS

    def size(self, run, para):
        el = self.rprop(run, para, "w:sz")
        return int(el.get(qn("w:val"))) / 2 if el is not None else 10.0

    def font(self, run, para):
        el = self.rprop(run, para, "w:rFonts")
        if el is None:
            return self._theme_fonts.get("minor", "Calibri")
        for a in ("w:ascii", "w:hAnsi"):
            if el.get(qn(a)):
                return el.get(qn(a))
        th = el.get(qn("w:asciiTheme")) or el.get(qn("w:hAnsiTheme")) or ""
        return self._theme_fonts.get("major" if th.startswith("major") else "minor", "Calibri")

    def color(self, run, para):
        el = self.rprop(run, para, "w:color")
        return el.get(qn("w:val")) if el is not None else None

    def underline(self, run, para):
        el = self.rprop(run, para, "w:u")
        return el is not None and el.get(qn("w:val")) not in ("none", None)

    def highlight(self, run, para):
        el = self.rprop(run, para, "w:highlight")
        if el is not None and el.get(qn("w:val")) not in ("none", None):
            return el.get(qn("w:val"))
        el = run._r.rPr.find(qn("w:shd")) if run._r.rPr is not None else None
        if el is not None and el.get(qn("w:fill")) not in (None, "auto", "FFFFFF"):
            return el.get(qn("w:fill"))
        return None

    # ---------- paragraph properties
    def alignment(self, para):
        v = self.pprop(para, "w:jc", "w:val")
        return ALIGN_NAME.get(v or "left", v)

    def line_spacing(self, para):
        line = self.pprop(para, "w:spacing", "w:line")
        rule = self.pprop(para, "w:spacing", "w:lineRule") or "auto"
        if line is None:
            return 1.0, "auto"
        if rule == "auto":
            return round(int(line) / 240, 2), "auto"
        return round(int(line) / 20, 1), rule        # points

    def indent(self, para):
        def tw(attr_list):
            for a in attr_list:
                v = self.pprop(para, "w:ind", a)
                if v is not None:
                    return int(v) / 1440
            return 0.0
        left = tw(["w:left", "w:start"])
        first = tw(["w:firstLine"]) - tw(["w:hanging"])
        right = tw(["w:right", "w:end"])
        return left, first, right

    def numbering(self, para):
        """Return the list format of a paragraph: 'bullet', 'decimal', 'lowerLetter', ... or None."""
        num_id, ilvl = None, "0"
        for ppr in self.ppr_chain(para):
            npr = ppr.find(qn("w:numPr"))
            if npr is not None and npr.find(qn("w:numId")) is not None:
                num_id = npr.find(qn("w:numId")).get(qn("w:val"))
                il = npr.find(qn("w:ilvl"))
                ilvl = il.get(qn("w:val")) if il is not None else "0"
                break
        else:
            return None
        if not num_id or num_id == "0":
            return None
        try:
            numbering = self.doc.part.numbering_part.element
        except Exception:
            return "list"
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id:
                aid = num.find(qn("w:abstractNumId")).get(qn("w:val"))
                for an in numbering.findall(qn("w:abstractNum")):
                    if an.get(qn("w:abstractNumId")) == aid:
                        for lvl in an.findall(qn("w:lvl")):
                            if lvl.get(qn("w:ilvl")) == ilvl:
                                f = lvl.find(qn("w:numFmt"))
                                return f.get(qn("w:val")) if f is not None else "list"
        return "list"

    # ---------- whole-document things
    def body_xml(self):
        return self.part_xml("word/document.xml")

    def header_footer_xml(self, kind):
        return "".join(self.part_xml(n) for n in self.parts(f"word/{kind}"))

    def header_footer_text(self, kind):
        out = []
        for n in self.parts(f"word/{kind}"):
            root = etree.fromstring(self._zip.read(n))
            out.append("".join(root.itertext()))
        return " ".join(out)

    def fields(self, xml=None):
        xml = xml if xml is not None else self.body_xml()
        out = re.findall(r"<w:instrText[^>]*>([^<]*)</w:instrText>", xml)
        out += re.findall(r'w:instr="([^"]*)"', xml)
        return [re.sub(r"\s+", " ", f).strip() for f in out]

    def sect_prs(self):
        return list(self.body.iter(qn("w:sectPr")))


def p_text(p):
    return "".join(Run(r, p).text for r in p._p.iter(qn("w:r")))


def norm_find(full, sub):
    """Find sub in full ignoring case/space differences; returns (start, end) in full."""
    low = full.lower()
    i = low.find(sub.lower())
    if i >= 0:
        return i, i + len(sub)
    pattern = r"\s+".join(re.escape(w) for w in sub.lower().split())
    m = re.search(pattern, low)
    if m:
        return m.start(), m.end()
    return None


# =================================================================== check builders

T_FMT = "Formatting documents"
T_PAGE = "Setting Page Style"
T_TABLE = "Creating Tables"
T_DRAW = "Drawing"
T_TOOLS = "Tools"
T_EDIT = "Editing text"


def _target(ctx, file, text):
    d = ctx.word(file)
    p = d.find_para(text)
    if p is None:
        return d, None, f"Text '{short(text, 50)}' was not found in the document."
    return d, p, None


def _runs_check(text, prop_ok, describe, file=None, sub=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        runs = d.runs_for(p, sub or text)
        if not runs:
            return False, "Text found, but it has no characters to check."
        bad = [r for r in runs if not prop_ok(d, r, p)]
        if bad:
            return False, describe(d, bad[0], p)
        return True, describe(d, runs[0], p)
    return fn


def typed(text, marks=2.0, file=None, min_acc=0.9, what=None):
    """The student typed the given passage (allowing a few typing mistakes)."""
    def fn(ctx):
        d = ctx.word(file)
        full = norm(d.text())
        want = norm(text)
        if want in full:
            return True, "Typed correctly (100% match)"
        # compare with the best window of the same length
        best = 0.0
        step = max(1, len(want) // 8)
        for i in range(0, max(1, len(full) - len(want) // 2), step):
            best = max(best, similarity(want, full[i:i + len(want) + 10]))
        ok = best >= min_acc
        return ok, f"Typing accuracy about {best:.0%}" + ("" if ok else f" (needs {min_acc:.0%} or more)")
    return Check(what or f"Typed the text: '{short(text, 45)}'", marks, fn,
                 expected="Text typed as given, with very few mistakes",
                 how="Type the passage carefully. Use Review > Spelling & Grammar (F7) to catch mistakes before saving.",
                 topic=T_EDIT)


def contains(text, marks=1.0, file=None, what=None, topic=T_EDIT, how="Type or insert the text as asked."):
    def fn(ctx):
        d = ctx.word(file)
        ok = norm(text) in norm(d.text())
        return ok, ("Found" if ok else f"'{short(text, 50)}' not found")
    return Check(what or f"Document contains '{short(text, 45)}'", marks, fn, expected=f"'{text}' present",
                 how=how, topic=topic)


def not_contains(words, marks=1.0, file=None, what=None, topic=T_EDIT, how=""):
    words = [words] if isinstance(words, str) else words

    def fn(ctx):
        d = ctx.word(file)
        full = d.text().lower()
        left = [w for w in words if re.search(r"\b" + re.escape(w.lower()) + r"\b", full)]
        return (not left), ("None left" if not left else "Still present: " + ", ".join(left))
    return Check(what or f"No '{', '.join(words)}' left in the document", marks, fn,
                 expected="These words should not appear any more", how=how, topic=topic)


def count_of(word, n, marks=1.0, file=None, what=None, topic="Find and Replace", how=""):
    def fn(ctx):
        d = ctx.word(file)
        c = len(re.findall(r"\b" + re.escape(word.lower()) + r"\b", d.text().lower()))
        return c >= n, f"'{word}' appears {c} time(s)"
    return Check(what or f"'{word}' appears {n} times", marks, fn, expected=f"{n} times", how=how, topic=topic)


def bold(text, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' is Bold", marks,
                 _runs_check(text, lambda d, r, p: d.is_on(r, p, "w:b"),
                             lambda d, r, p: "Bold" if d.is_on(r, p, "w:b") else f"'{short(r.text, 30)}' is not bold",
                             file, sub),
                 expected="Bold", how="Select the text > press Ctrl+B (or Home > B).", topic=T_FMT)


def italic(text, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' is Italic", marks,
                 _runs_check(text, lambda d, r, p: d.is_on(r, p, "w:i"),
                             lambda d, r, p: "Italic" if d.is_on(r, p, "w:i") else f"'{short(r.text, 30)}' is not italic",
                             file, sub),
                 expected="Italic", how="Select the text > press Ctrl+I (or Home > I).", topic=T_FMT)


def underline(text, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' is Underlined", marks,
                 _runs_check(text, lambda d, r, p: d.underline(r, p),
                             lambda d, r, p: "Underlined" if d.underline(r, p) else f"'{short(r.text, 30)}' is not underlined",
                             file, sub),
                 expected="Underlined", how="Select the text > press Ctrl+U (or Home > U).", topic=T_FMT)


def font_size(text, pt, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' font size is {pt} pt", marks,
                 _runs_check(text, lambda d, r, p: abs(d.size(r, p) - pt) < 0.1,
                             lambda d, r, p: f"Font size {d.size(r, p):g} pt", file, sub),
                 expected=f"{pt} pt", how=f"Select the text > Home > Font Size box > type {pt} > Enter.", topic=T_FMT)


def font_name(text, name, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' font is {name}", marks,
                 _runs_check(text, lambda d, r, p: d.font(r, p).lower() == name.lower(),
                             lambda d, r, p: f"Font {d.font(r, p)}", file, sub),
                 expected=name, how=f"Select the text > Home > Font box > choose {name}.", topic=T_FMT)


def font_color(text, color="any", marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' font colour is {color if color != 'any' else 'changed'}", marks,
                 _runs_check(text, lambda d, r, p: color_matches(d.color(r, p), color),
                             lambda d, r, p: f"Colour {d.color(r, p) or 'Automatic (black)'}", file, sub),
                 expected=f"{color.capitalize()} text" if color != "any" else "Any colour except black",
                 how="Select the text > Home > Font Color (the 'A' with a coloured bar) > pick the colour.", topic=T_FMT)


def highlighted(text, marks=1.0, file=None, sub=None):
    return Check(f"'{short(sub or text, 40)}' is highlighted", marks,
                 _runs_check(text, lambda d, r, p: bool(d.highlight(r, p)),
                             lambda d, r, p: f"Highlight: {d.highlight(r, p) or 'none'}", file, sub),
                 expected="Text highlight colour applied",
                 how="Select the text > Home > Text Highlight Color (the marker pen) > pick Yellow.", topic=T_FMT)


def uppercase(text, marks=1.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        for p in d.paragraphs():
            t = p_text(p)
            if norm(text) in norm(t):
                runs = d.runs_for(p, text)
                if all(d.is_on(r, p, "w:caps") for r in runs):
                    return True, "ALL CAPS applied"
                if text.upper() in t:
                    return True, "Text is in UPPERCASE"
                return False, "Text is not in uppercase"
        if text.upper() in d.text():
            return True, "Text is in UPPERCASE"
        return False, f"'{short(text, 40)}' not found"
    return Check(f"'{short(text, 40)}' changed to UPPERCASE", marks, fn, expected=text.upper(),
                 how="Select the text > Home > Change Case (Aa) > UPPERCASE. (Shift+F3 also cycles the case.)",
                 topic="Case settings")


def alignment(text, align, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        a = d.alignment(p)
        return a == align, f"Alignment: {a}"
    keys = {"Center": "Ctrl+E", "Justify": "Ctrl+J", "Right": "Ctrl+R", "Left": "Ctrl+L"}
    return Check(f"'{short(text, 40)}' is {align} aligned", marks, fn, expected=align,
                 how=f"Click in the paragraph > press {keys[align]} (or Home > Paragraph > {align}).",
                 topic="Alignments")


def line_spacing(text, value, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        v, rule = d.line_spacing(p)
        if rule != "auto":
            return False, f"Line spacing is set to {v} pt ({rule}), not multiple"
        return abs(v - value) < 0.06, f"Line spacing {v:g}"
    return Check(f"Line spacing of '{short(text, 35)}' is {value}", marks, fn, expected=f"{value} lines",
                 how=f"Click in the paragraph > Home > Line and Paragraph Spacing > {value}. "
                     f"(Or Paragraph dialog > Line spacing.) Shortcut for 1.5 is Ctrl+5.",
                 topic="Line Space")


def indent(text, left=None, first=None, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        l, f, _ = d.indent(p)
        ok = (left is None or abs(l - left) < 0.06) and (first is None or abs(f - first) < 0.06)
        return ok, f"Left indent {l:.2f}\", first line {f:+.2f}\""
    parts = []
    if left is not None:
        parts.append(f"left indent {left}\"")
    if first is not None:
        parts.append(f"first line indent {first}\"")
    return Check(f"'{short(text, 35)}' has " + " and ".join(parts), marks, fn, expected=", ".join(parts),
                 how="Click in the paragraph > Home > Paragraph group dialog launcher (small arrow) > Indentation: "
                     "set Left and Special = First line by the values asked > OK.", topic="Indents")


def spacing_after(text, pt, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        v = d.pprop(p, "w:spacing", "w:after")
        v = int(v) / 20 if v is not None else 0
        return abs(v - pt) < 0.6, f"Spacing after {v:g} pt"
    return Check(f"Spacing after '{short(text, 35)}' is {pt} pt", marks, fn, expected=f"{pt} pt",
                 how="Click in the paragraph > Layout (Page Layout) tab > Spacing After box > type the value.",
                 topic="Line Space")


def style(text, name, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        s = d.style_name(p)
        return s.lower() == name.lower(), f"Style: {s}"
    return Check(f"'{short(text, 40)}' uses style {name}", marks, fn, expected=name,
                 how=f"Click in the line > Home > Styles gallery > click '{name}'.", topic="Setting paragraph style")


def listed(texts, kind="bullet", marks=1.0, file=None):
    """kind: 'bullet' or 'number' (any numbering format)."""
    def fn(ctx):
        d = ctx.word(file)
        bad = []
        for t in texts:
            p = d.find_para(t)
            if p is None:
                bad.append(f"'{short(t, 25)}' not found")
                continue
            f = d.numbering(p)
            if f is None:
                bad.append(f"'{short(t, 25)}' has no {kind}")
            elif kind == "bullet" and f != "bullet":
                bad.append(f"'{short(t, 25)}' is numbered, not bulleted")
            elif kind == "number" and f in ("bullet", "none"):
                bad.append(f"'{short(t, 25)}' has bullets, not numbers")
        return (not bad), ("All items OK" if not bad else "; ".join(bad))
    label = "bullets" if kind == "bullet" else "numbering"
    return Check(f"{len(texts)} items have {label}", marks, fn, expected=f"Each item starts with a {kind}",
                 how=f"Select the lines > Home > {'Bullets' if kind == 'bullet' else 'Numbering'} button.",
                 topic="Bullets & Numbering")


# ---------- page setup

PAPER = {"A4": (8.27, 11.69), "Letter": (8.5, 11.0), "Legal": (8.5, 14.0)}


def orientation(kind="landscape", marks=1.0, file=None):
    def fn(ctx):
        s = ctx.word(file).doc.sections[0]
        land = s.page_width > s.page_height
        got = "Landscape" if land else "Portrait"
        return got.lower() == kind, f"Orientation: {got}"
    return Check(f"Page orientation is {kind.capitalize()}", marks, fn, expected=kind.capitalize(),
                 how=f"Layout (Page Layout) tab > Orientation > {kind.capitalize()}.", topic=T_PAGE)


def paper_size(name="A4", marks=1.0, file=None):
    def fn(ctx):
        s = ctx.word(file).doc.sections[0]
        w, h = sorted([s.page_width.inches, s.page_height.inches])
        pw, ph = PAPER[name]
        got = next((k for k, (a, b) in PAPER.items() if abs(a - w) < 0.05 and abs(b - h) < 0.05), f"{w:.2f}\" x {h:.2f}\"")
        return got == name, f"Paper size: {got}"
    return Check(f"Paper size is {name}", marks, fn, expected=name,
                 how=f"Layout (Page Layout) tab > Size > {name}.", topic=T_PAGE)


def margins(top=None, bottom=None, left=None, right=None, preset=None, marks=1.0, file=None):
    presets = {"Narrow": (0.5, 0.5, 0.5, 0.5), "Normal": (1, 1, 1, 1), "Moderate": (1, 1, 0.75, 0.75),
               "Wide": (1, 1, 2, 2)}
    if preset:
        top, bottom, left, right = presets[preset]

    def fn(ctx):
        s = ctx.word(file).doc.sections[0]
        got = [s.top_margin.inches, s.bottom_margin.inches, s.left_margin.inches, s.right_margin.inches]
        want = [top, bottom, left, right]
        ok = all(w is None or abs(g - w) < 0.03 for g, w in zip(got, want))
        return ok, "Margins T/B/L/R: " + " / ".join(f"{g:.2f}\"" for g in got)
    desc = f"Margins set to {preset}" if preset else "Margins T/B/L/R " + " / ".join(
        f"{v}\"" if v is not None else "-" for v in (top, bottom, left, right))
    how = (f"Layout (Page Layout) tab > Margins > {preset}." if preset else
           "Layout (Page Layout) tab > Margins > Custom Margins > type Top, Bottom, Left, Right > OK. "
           "(If your Word shows cm, 1 inch = 2.54 cm.)")
    return Check(desc, marks, fn, expected=desc.replace("Margins set to ", ""), how=how, topic=T_PAGE)


def columns(n=2, marks=1.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        nums = []
        for sp in d.sect_prs():
            c = sp.find(qn("w:cols"))
            nums.append(int(c.get(qn("w:num")) or 1) if c is not None else 1)
        return (n in nums), f"Column settings found: {sorted(set(nums))}"
    return Check(f"Text set in {n} columns", marks, fn, expected=f"{n} columns",
                 how=f"Select the paragraph(s) > Layout (Page Layout) > Columns > {'Two' if n == 2 else 'Three'}.",
                 topic="Column break")


def page_border(marks=1.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        ok = any(sp.find(qn("w:pgBorders")) is not None for sp in d.sect_prs())
        return ok, "Page border present" if ok else "No page border"
    return Check("Page border applied", marks, fn, expected="A border around the page",
                 how="Design tab (Word 2013+) or Page Layout tab (Word 2010) > Page Borders > Box > choose style > OK.",
                 topic="Border & Shading")


def para_border_or_shading(text, marks=1.0, file=None):
    def fn(ctx):
        d, p, err = _target(ctx, file, text)
        if err:
            return False, err
        ppr = p._p.pPr
        b = ppr is not None and ppr.find(qn("w:pBdr")) is not None
        s = ppr is not None and ppr.find(qn("w:shd")) is not None and ppr.find(qn("w:shd")).get(qn("w:fill")) not in (None, "auto", "FFFFFF")
        if not s:
            s = any(r._r.rPr is not None and r._r.rPr.find(qn("w:shd")) is not None for r in d.runs_for(p))
        if not b:
            b = any(r._r.rPr is not None and r._r.rPr.find(qn("w:bdr")) is not None for r in d.runs_for(p))
        return (b and s), f"Border: {'yes' if b else 'no'}, Shading: {'yes' if s else 'no'}"
    return Check(f"Border and shading on '{short(text, 35)}'", marks, fn, expected="Both a border and a shading colour",
                 how="Select the paragraph > Home > Borders drop-down > Borders and Shading > Borders tab: Box; "
                     "Shading tab: pick a Fill colour > Apply to: Paragraph > OK.", topic="Border & Shading")


def page_color(marks=1.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        bg = d.doc.element.find(qn("w:background"))
        ok = bg is not None and (bg.get(qn("w:color")) not in (None, "FFFFFF", "auto") or len(bg))
        return ok, "Page colour set" if ok else "No page colour"
    return Check("Page colour applied", marks, fn, expected="Page background colour",
                 how="Design tab > Page Color (Word 2010: Page Layout > Page Color) > pick a light colour.",
                 topic=T_PAGE)


def header_text(text, marks=1.0, file=None, kind="header"):
    def fn(ctx):
        d = ctx.word(file)
        t = d.header_footer_text(kind)
        ok = norm(text) in norm(t)
        return ok, f"{kind.capitalize()} text: '{short(t, 60)}'" if t.strip() else f"No {kind} text"
    return Check(f"{kind.capitalize()} contains '{short(text, 40)}'", marks, fn, expected=text,
                 how=f"Insert > {kind.capitalize()} > Edit {kind.capitalize()} (or double-click the top/bottom of the page) "
                     f"> type the text > Close {kind.capitalize()} and Footer.", topic="Header & Footer")


def footer_text(text, marks=1.0, file=None):
    return header_text(text, marks, file, kind="footer")


def page_numbers(marks=1.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        xml = d.header_footer_xml("header") + d.header_footer_xml("footer")
        ok = any(re.match(r"PAGE\b", f) for f in d.fields(xml))
        return ok, "Page number field found" if ok else "No page number in header/footer"
    return Check("Page numbers inserted", marks, fn, expected="Page number in header or footer",
                 how="Insert > Page Number > Bottom of Page > Plain Number 2 (centre).", topic="Header & Footer")


def page_break(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ok = 'w:type="page"' in x or "<w:pageBreakBefore" in x
        return ok, "Page break found" if ok else "No manual page break"
    return Check("Manual page break inserted", marks, fn, expected="A page break",
                 how="Click where the new page should start > press Ctrl+Enter (or Insert > Page Break).",
                 topic="Column break and line break")


def column_break(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ok = 'w:type="column"' in x
        return ok, "Column break found" if ok else "No column break"
    return Check("Column break inserted", marks, fn, expected="A column break",
                 how="Click where the next column should start > Layout (Page Layout) > Breaks > Column "
                     "(or Ctrl+Shift+Enter).", topic="Column break and line break")


def line_break(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        n = len(re.findall(r"<w:br(?![^>]*w:type=\"(?:page|column)\")[^>]*/>", x))
        return n > 0, f"{n} line break(s) found"
    return Check("Line break (Shift+Enter) used", marks, fn, expected="A manual line break",
                 how="Click at the place > press Shift+Enter.", topic="Column break and line break")


def symbol(chars, marks=1.0, file=None):
    chars = list(chars)

    def fn(ctx):
        d = ctx.word(file)
        t = d.text()
        x = d.body_xml()
        found = [c for c in chars if c in t]
        if found:
            return True, f"Found symbol(s): {' '.join(found)}"
        if "<w:sym " in x:
            return True, "Symbol inserted (from a symbol font)"
        return False, "Symbol not found"
    return Check(f"Symbol inserted ({' or '.join(chars)})", marks, fn, expected=" or ".join(chars),
                 how="Insert > Symbol > More Symbols > pick the symbol > Insert > Close. "
                     "(© is Alt+Ctrl+C, ® is Alt+Ctrl+R, ™ is Alt+Ctrl+T.)", topic="Special symbols")


def picture(n=1, marks=1.0, file=None, where="body"):
    def fn(ctx):
        d = ctx.word(file)
        x = d.body_xml() if where == "body" else d.header_footer_xml("header")
        c = x.count("<pic:pic") + x.count("<v:imagedata")
        return c >= n, f"{c} picture(s) found"
    return Check(f"{n} picture/clip art inserted" + (" in header" if where != "body" else ""), marks, fn,
                 expected=f"At least {n} picture",
                 how="Insert > Pictures (This Device) > choose an image (the test folder has sample pictures) > Insert. "
                     "Or Insert > Online Pictures / Clip Art in older Word.", topic=T_DRAW)


def shape(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        c = x.count("<wps:wsp") or len(re.findall(r"<v:(shape|rect|oval|roundrect)\b", x))
        return c > 0, f"{c} shape(s) found"
    return Check("Shape drawn", marks, fn, expected="At least one shape",
                 how="Insert > Shapes > pick a shape > drag on the page.", topic=T_DRAW)


def wordart(text, marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        boxes = re.findall(r"<w:txbxContent>(.*?)</w:txbxContent>", x, re.S) + re.findall(r"<v:textpath[^>]*string=\"([^\"]*)\"", x)
        joined = " ".join(re.sub(r"<[^>]+>", "", b) for b in boxes)
        ok = norm(text) in norm(joined)
        return ok, "WordArt/text box with the text found" if ok else "No WordArt/text box with that text"
    return Check(f"WordArt with '{text}'", marks, fn, expected=f"WordArt reading '{text}'",
                 how="Insert > WordArt > pick a style > type the text.", topic=T_DRAW)


def hyperlink(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ok = "<w:hyperlink" in x or any(f.startswith("HYPERLINK") for f in ctx.word(file).fields())
        return ok, "Hyperlink found" if ok else "No hyperlink"
    return Check("Hyperlink inserted", marks, fn, expected="A clickable link",
                 how="Select the text > Ctrl+K (Insert > Link) > type the address > OK.", topic="Inserting")


def drop_cap(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ok = "w:dropCap=" in x
        return ok, "Drop cap found" if ok else "No drop cap"
    return Check("Drop cap applied to first paragraph", marks, fn, expected="Large first letter",
                 how="Click in the first paragraph > Insert > Drop Cap > Dropped.", topic=T_FMT)


def footnote(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ok = "<w:footnoteReference" in x
        return ok, "Footnote found" if ok else "No footnote"
    return Check("Footnote inserted", marks, fn, expected="A footnote",
                 how="Click after the word > References > Insert Footnote (Alt+Ctrl+F) > type the note.",
                 topic="Inserting")


# ---------- tables

def _table(ctx, file, index):
    d = ctx.word(file)
    tables = d.doc.tables
    if len(tables) <= index:
        return d, None, f"Only {len(tables)} table(s) found"
    return d, tables[index], None


def _grid(tbl):
    rows = []
    for tr in tbl._tbl.findall(qn("w:tr")):
        cells = []
        for tc in tr.findall(qn("w:tc")):
            cells.append(" ".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in tc.iter(qn("w:p"))).strip())
        rows.append(cells)
    return rows


def table_size(rows, cols, marks=1.0, file=None, index=0, at_least=False):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        r = len(t._tbl.findall(qn("w:tr")))
        g = t._tbl.find(qn("w:tblGrid"))
        c = len(g.findall(qn("w:gridCol"))) if g is not None else len(t.columns)
        ok = (r >= rows and c >= cols) if at_least else (r == rows and c == cols)
        return ok, f"Table has {r} rows x {c} columns"
    return Check(f"Table of {rows} rows x {cols} columns", marks, fn, expected=f"{rows} x {cols}",
                 how=f"Insert > Table > Insert Table > Columns {cols}, Rows {rows} > OK.", topic=T_TABLE)


def table_has(texts, marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        flat = norm(" | ".join(" | ".join(r) for r in _grid(t)))
        miss = [x for x in texts if norm(x) not in flat]
        return (not miss), ("All entries present" if not miss else "Missing: " + ", ".join(miss[:5]))
    return Check("Table data entered correctly", marks, fn, expected="All the given values in the table",
                 how="Click in each cell and type; use Tab to move to the next cell.", topic=T_TABLE)


def table_merged(marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        x = etree.tostring(t._tbl).decode()
        ok = "<w:gridSpan" in x or "<w:vMerge" in x
        return ok, "Merged cells found" if ok else "No merged cells"
    return Check("Cells merged in the table", marks, fn, expected="Merged cells",
                 how="Select the cells > Table Tools Layout > Merge Cells (or right-click > Merge Cells).",
                 topic=T_TABLE)


def table_borders(marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        x = etree.tostring(t._tbl).decode()
        style = t._tbl.tblPr.find(qn("w:tblStyle"))
        sid = style.get(qn("w:val")) if style is not None else ""
        own = re.search(r"<w:(tblBorders|tcBorders)>.*?w:val=\"(single|double|thick|dotted|dashed)", x, re.S)
        ok = bool(own) or (sid and sid not in ("TableNormal", "NormalTable"))
        return ok, f"Table style: {sid or 'none'}; own borders: {'yes' if own else 'no'}"
    return Check("Table has borders", marks, fn, expected="Visible borders on the table",
                 how="Select the table > Table Tools Design > Borders > All Borders.", topic="Borders")


def table_shading(marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        x = etree.tostring(t._tbl).decode()
        fills = [f for f in re.findall(r"<w:shd [^>]*w:fill=\"([0-9A-Fa-f]{6}|auto)\"", x) if f.lower() not in ("auto", "ffffff")]
        style = t._tbl.tblPr.find(qn("w:tblStyle"))
        sid = style.get(qn("w:val")) if style is not None else ""
        styled = sid and sid not in ("TableNormal", "NormalTable", "TableGrid")
        ok = bool(fills) or bool(styled)
        return ok, f"Shaded cells: {len(fills)}; table style: {sid or 'none'}"
    return Check("Shading applied in the table", marks, fn, expected="Header row (or cells) shaded",
                 how="Select the header row > Table Tools Design > Shading > pick a colour.", topic="Shading")


def table_formula(func="SUM", expect=None, marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        x = etree.tostring(t._tbl).decode()
        flds = d.fields(x)
        f = [f for f in flds if f.upper().startswith("=" + func.upper()) or f.upper().startswith(f"={func.upper()}(")]
        if not f:
            return False, "No table formula found (the total may have been typed by hand)"
        if expect is not None:
            flat = " ".join(" ".join(r) for r in _grid(t)).replace(",", "")
            if str(expect) not in flat:
                return False, f"Formula {f[0]} found but the result {expect} is not shown"
        return True, f"Formula {f[0]}"
    return Check(f"Total calculated with a table formula (={func}(...))", marks, fn,
                 expected=f"Field ={func}(ABOVE)" + (f" showing {expect}" if expect is not None else ""),
                 how=f"Click in the total cell > Table Tools Layout > Formula > keep =SUM(ABOVE) > OK.",
                 topic="Formula")


def table_sorted(col, desc=False, numeric=False, marks=1.0, file=None, index=0):
    def fn(ctx):
        d, t, err = _table(ctx, file, index)
        if err:
            return False, err
        rows = [r for r in _grid(t)[1:] if not any(c.lower().startswith("total") for c in r)]
        vals = [r[col] for r in rows if len(r) > col and r[col]]
        key = (lambda v: float(re.sub(r"[^0-9.\-]", "", v) or 0)) if numeric else (lambda v: v.lower())
        ok = len(vals) > 1 and [key(v) for v in vals] == sorted([key(v) for v in vals], reverse=desc)
        return ok, "Order: " + short(", ".join(vals), 70)
    return Check(f"Table sorted by column {col + 1} ({'descending' if desc else 'ascending'})", marks, fn,
                 expected=("Largest first" if desc else "A to Z / smallest first"),
                 how="Click in the table > Table Tools Layout > Sort > Sort by: the column; choose "
                     f"{'Descending' if desc else 'Ascending'}; My list has: Header row > OK.", topic="Sorting")


def table_inserted_row(min_rows, marks=1.0, file=None, index=0):
    return table_size(min_rows, 1, marks, file, index, at_least=True)


# ---------- tools

def track_changes(marks=1.0, file=None):
    def fn(ctx):
        x = ctx.word(file).body_xml()
        ins, dl = x.count("<w:ins "), x.count("<w:del ")
        return (ins + dl) > 0, f"{ins} tracked insertion(s), {dl} tracked deletion(s)"
    return Check("Edits made with Track Changes ON", marks, fn, expected="Tracked insertions/deletions visible",
                 how="Review > Track Changes (Ctrl+Shift+E) to turn it on, THEN make the edits, then save.",
                 topic="Tracking changes")


def comment(marks=1.0, file=None, text=None):
    def fn(ctx):
        d = ctx.word(file)
        x = d.part_xml("word/comments.xml")
        if "<w:comment " not in x:
            return False, "No comment found"
        if text and norm(text) not in norm(re.sub(r"<[^>]+>", " ", x)):
            return False, "Comment found but not with the asked text"
        return True, "Comment found"
    return Check("Comment inserted" + (f": '{text}'" if text else ""), marks, fn, expected="A comment",
                 how="Select the word > Review > New Comment > type the comment.", topic="Tracking changes")


def mergefields(names, marks=2.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        flds = " ".join(d.fields()).upper()
        miss = [n for n in names if f"MERGEFIELD {n.upper()}" not in flds.replace('"', "").replace("  ", " ")]
        return (not miss), ("All merge fields present" if not miss else "Missing merge field(s): " + ", ".join(miss))
    return Check("Merge fields inserted: " + ", ".join(f"«{n}»" for n in names), marks, fn,
                 expected="«" + "», «".join(names) + "» in the letter",
                 how="Mailings > Start Mail Merge > Letters > Select Recipients > Use an Existing List > pick the Excel "
                     "file > OK. Then put the cursor in place > Insert Merge Field > choose each field.",
                 topic="Mail Merge")


def merged_output(values, marks=2.0, file=None):
    def fn(ctx):
        d = ctx.word(file)
        t = norm(d.text())
        miss = [v for v in values if norm(v) not in t]
        if "mergefield" in " ".join(d.fields()).lower():
            return False, "This file still has merge fields. Use Finish & Merge > Edit Individual Documents."
        return (not miss), (f"All {len(values)} letters generated" if not miss else "Missing: " + ", ".join(miss))
    return Check("Merged letters generated for all recipients", marks, fn,
                 expected="One letter per recipient with real names",
                 how="Mailings > Finish & Merge > Edit Individual Documents > All > OK. Save the new document.",
                 topic="Mail Merge")


def spell_fixed(pairs, marks=2.0, file=None):
    """pairs = {wrong: right}"""
    def fn(ctx):
        d = ctx.word(file)
        t = d.text().lower()
        left = [w for w in pairs if re.search(r"\b" + w.lower() + r"\b", t)]
        fixed = [r for r in pairs.values() if r.lower() in t]
        ok = not left and len(fixed) == len(pairs)
        msg = f"{len(fixed)} of {len(pairs)} words corrected"
        if left:
            msg += "; still wrong: " + ", ".join(left)
        return ok, msg
    return Check("Spelling mistakes corrected", marks, fn,
                 expected="; ".join(f"{a} → {b}" for a, b in pairs.items()),
                 how="Press F7 (Review > Spelling & Grammar) and click Change for each wrong word.", topic="Spell Checks")
