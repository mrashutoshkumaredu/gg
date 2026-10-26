"""Reading an .xlsx and the MS Excel check builders."""
from __future__ import annotations

import re
import zipfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from .core import Check, Missing, norm, partial, short

T_FORMULA = "Setting formula"
T_FMT = "Formatting spreadsheets"
T_SHEETS = "Working with sheets"
T_DATA = "Entering & Deleting Data"


class Book:
    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(str(path), data_only=False)
        self.wbv = load_workbook(str(path), data_only=True)
        self.zip = zipfile.ZipFile(str(path))

    def ws(self, name=None, values=False):
        wb = self.wbv if values else self.wb
        if name is None:
            return wb.worksheets[0]
        for alt in name.split("|"):              # "New Name|Old Name": still check the work if not renamed
            for s in wb.worksheets:
                if s.title.lower() == alt.lower():
                    return s
        raise Missing(f"Sheet '{name.split('|')[0]}' not found (sheets are: {', '.join(wb.sheetnames)})")

    def xml(self, prefix):
        return "".join(self.zip.read(n).decode("utf8", "ignore") for n in self.zip.namelist() if n.startswith(prefix))

    def sheet_xml(self, name):
        """Raw XML of one worksheet (to find things openpyxl drops)."""
        wbx = self.zip.read("xl/workbook.xml").decode("utf8", "ignore")
        rels = self.zip.read("xl/_rels/workbook.xml.rels").decode("utf8", "ignore")
        for m in re.finditer(r"<sheet [^>]*>", wbx):
            tag = m.group(0)
            nm = re.search(r'name="([^"]*)"', tag).group(1)
            rid = re.search(r'r:id="([^"]*)"', tag).group(1)
            if name is None or nm.lower() in [a.lower() for a in name.split("|")]:
                t = re.search(r'Id="%s"[^>]*Target="([^"]*)"' % rid, rels) or re.search(r'Target="([^"]*)"[^>]*Id="%s"' % rid, rels)
                target = t.group(1).lstrip("/")
                if not target.startswith("xl/"):
                    target = "xl/" + target
                return self.zip.read(target).decode("utf8", "ignore")
        return ""


def cells(ws, rng):
    for row in ws[rng] if ":" in rng else [[ws[rng]]]:
        for c in row:
            yield c


def _all_empty(ws, rng):
    return all(c.value is None for c in cells(ws, rng))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _sheet_label(sheet):
    return f"{sheet.split('|')[0]}!" if sheet else ""


# =================================================================== data / formula checks

def value(sheet, ref, expect, marks=1.0, file=None, what=None, tol=0.01, topic=T_DATA):
    def fn(ctx):
        b = ctx.excel(file)
        v = b.ws(sheet, values=True)[ref].value
        if v is None:
            v = b.ws(sheet)[ref].value
        if isinstance(expect, (int, float)):
            n = _num(v)
            return n is not None and abs(n - expect) <= tol, f"{ref} = {v}"
        return norm(str(v)) == norm(str(expect)), f"{ref} = {v}"
    return Check(what or f"{_sheet_label(sheet)}{ref} contains {expect}", marks, fn, expected=str(expect),
                 how=f"Click {ref} and type the value.", topic=topic)


def formula(sheet, ref, func, expect=None, marks=1.0, file=None, what=None, how=None, tol=0.01, topic=T_FORMULA):
    """func: a function name like SUM, or a regex the formula must match."""
    def fn(ctx):
        b = ctx.excel(file)
        f = b.ws(sheet)[ref].value
        v = b.ws(sheet, values=True)[ref].value
        if f is None:
            return False, f"{ref} is empty"
        if not (isinstance(f, str) and f.startswith("=")):
            return False, f"{ref} has a typed value ({f}), not a formula"
        if not re.search(func if any(ch in func for ch in "\\[(*+$") else r"\b" + func + r"\s*\(", f, re.I):
            return False, f"{ref} has {f}, which does not use {func}"
        if expect is not None and v is not None:
            n = _num(v)
            if isinstance(expect, (int, float)):
                if n is None or abs(n - expect) > tol:
                    return False, f"{ref} has {f} but the answer is {v} (should be {expect})"
            elif norm(str(v)) != norm(str(expect)):
                return False, f"{ref} has {f} but the answer is {v} (should be {expect})"
        return True, f"{ref}: {f}" + (f" = {v}" if v is not None else "")
    return Check(what or f"{_sheet_label(sheet)}{ref} uses a {func} formula", marks, fn,
                 expected=(f"={func}(...)" if func.isalpha() else "A formula") + (f" giving {expect}" if expect is not None else ""),
                 how=how or f"Click {ref} > type ={func}( and select the range > ) > Enter. Or Formulas > Insert Function.",
                 topic=topic)


def series(sheet, rng, start, step=1, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet, values=True)
        vals = [c.value for c in cells(ws, rng)]
        want = [start + i * step for i in range(len(vals))]
        ok = all(_num(v) == w for v, w in zip(vals, want))
        return ok, "Values: " + short(", ".join(str(v) for v in vals), 60)
    return Check(f"{rng} filled with series {start}, {start + step}, …", marks, fn, expected="A continuous series",
                 how=f"Type {start} and {start + step} in the first two cells, select both, drag the fill handle down. "
                     "Or Home > Fill > Series.", topic="Filling continuous rows/columns")


def text_in(sheet, rng, texts, marks=1.0, file=None, what=None, topic=T_DATA):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet, values=True)
        have = norm(" | ".join(str(c.value) for c in cells(ws, rng) if c.value is not None))
        miss = [t for t in texts if norm(str(t)) not in have]
        return partial(len(texts) - len(miss), len(texts)), ("All present" if not miss else "Missing: " + ", ".join(map(str, miss[:5])))
    return Check(what or f"Data entered in {_sheet_label(sheet)}{rng}", marks, fn, expected=", ".join(map(str, texts[:6])),
                 how="Click each cell and type; press Enter or Tab to move on.", topic=topic)


def no_errors(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet, values=True)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        errs = [f"{c.coordinate}={c.value}" for c in cells(ws, rng) if isinstance(c.value, str) and c.value.startswith("#")]
        return (not errs), ("No error values" if not errs else "Errors: " + ", ".join(errs[:5]))
    return Check(f"No error values (#DIV/0!, #N/A…) in {rng}", marks, fn, expected="No # errors",
                 how="Wrap the formula with IFERROR, e.g. =IFERROR(B2/C2,0). Formulas > Error Checking shows the cells.",
                 topic="Error checking")


# =================================================================== formatting

def bold(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [c.coordinate for c in cells(ws, rng) if c.value is not None and not c.font.b]
        return (not bad), ("All bold" if not bad else "Not bold: " + ", ".join(bad[:5]))
    return Check(f"{_sheet_label(sheet)}{rng} is Bold", marks, fn, expected="Bold",
                 how="Select the cells > Ctrl+B.", topic=T_FMT)


def font_size(sheet, rng, pt, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [f"{c.coordinate}={c.font.sz}" for c in cells(ws, rng) if c.value is not None and float(c.font.sz or 11) != pt]
        return (not bad), ("All OK" if not bad else "Size: " + ", ".join(bad[:4]))
    return Check(f"{_sheet_label(sheet)}{rng} font size {pt}", marks, fn, expected=f"{pt} pt",
                 how=f"Select the cells > Home > Font Size > {pt}.", topic="font")


def font_color(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        def coloured(c):
            col = c.font.color
            return col is not None and ((col.rgb and isinstance(col.rgb, str) and col.rgb[-6:] not in ("000000",)) or col.theme not in (None, 1))
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [c.coordinate for c in cells(ws, rng) if c.value is not None and not coloured(c)]
        return (not bad), ("Coloured" if not bad else "Not coloured: " + ", ".join(bad[:5]))
    return Check(f"{_sheet_label(sheet)}{rng} font colour changed", marks, fn, expected="Font colour other than black",
                 how="Select > Home > Font Color > pick a colour.", topic="font")


def fill(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        def filled(c):
            f = c.fill
            if f is None:
                return False
            if getattr(f, "type", None) in ("linear", "path"):      # gradient fill
                return True
            return getattr(f, "fill_type", None) not in (None, "none")
        bad = [c.coordinate for c in cells(ws, rng) if not filled(c)]
        return (not bad), ("Shaded" if not bad else "No fill: " + ", ".join(bad[:5]))
    return Check(f"{_sheet_label(sheet)}{rng} has a fill (shading) colour", marks, fn, expected="Shaded cells",
                 how="Select > Home > Fill Color (paint bucket) > pick a colour.", topic="border & shading")


def borders(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        min_c, min_r, max_c, max_r = range_boundaries(rng)
        bad = []
        for row in range(min_r, max_r + 1):
            for col in range(min_c, max_c + 1):
                c = ws.cell(row, col)
                bd = c.border
                # an edge counts if this cell or its neighbour draws it
                left = bd.left.style or (col > 1 and ws.cell(row, col - 1).border.right.style)
                right = bd.right.style or ws.cell(row, col + 1).border.left.style
                top = bd.top.style or (row > 1 and ws.cell(row - 1, col).border.bottom.style)
                bottom = bd.bottom.style or ws.cell(row + 1, col).border.top.style
                if not (left and right and top and bottom):
                    bad.append(c.coordinate)
        total = (max_r - min_r + 1) * (max_c - min_c + 1)
        return (1.0 if not bad else partial(total - len(bad), total) * 0.75), ("All borders present" if not bad else f"{len(bad)} cell(s) without full borders, e.g. {', '.join(bad[:4])}")
    return Check(f"All borders on {_sheet_label(sheet)}{rng}", marks, fn, expected="Every cell has a border",
                 how="Select the range > Home > Borders drop-down > All Borders.", topic="border & shading")


def merged(sheet, rng, marks=1.0, file=None, centre=True):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        mr = [str(m) for m in ws.merged_cells.ranges]
        a = rng.split(":")[0].upper()
        hit = [m for m in mr if m.upper().startswith(a + ":")]
        if not hit:
            return False, "Not merged" + (f" (merged ranges: {', '.join(mr)})" if mr else "")
        if hit[0].upper() != rng.upper():
            return False, f"Merged {hit[0]} instead of {rng}"
        if centre and ws[a].alignment.horizontal not in ("center", "centerContinuous"):
            return False, f"Merged {hit[0]} but not centred"
        return True, f"Merged {hit[0]}" + (" and centred" if centre else "")
    return Check(f"{_sheet_label(sheet)}{rng} merged" + (" and centred" if centre else ""), marks, fn,
                 expected=f"{rng} as one cell" + (", text in the centre" if centre else ""),
                 how=f"Select {rng} > Home > Merge & Center.", topic="Alignment")


def align(sheet, rng, how_="center", marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [c.coordinate for c in cells(ws, rng) if c.value is not None and (c.alignment.horizontal or "general") != how_]
        return (not bad), ("Aligned" if not bad else "Not aligned: " + ", ".join(bad[:5]))
    return Check(f"{_sheet_label(sheet)}{rng} aligned {how_}", marks, fn, expected=how_.capitalize(),
                 how=f"Select > Home > Alignment > {'Center' if how_ == 'center' else how_.capitalize()}.", topic="Alignment")


def wrap(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [c.coordinate for c in cells(ws, rng) if c.value is not None and not c.alignment.wrap_text]
        return (not bad), ("Wrap text on" if not bad else "Not wrapped: " + ", ".join(bad[:5]))
    return Check(f"Wrap Text on {_sheet_label(sheet)}{rng}", marks, fn, expected="Wrap Text",
                 how="Select > Home > Wrap Text.", topic="Alignment")


def number_format(sheet, rng, kind, marks=1.0, file=None):
    tests = {
        "currency": lambda f: any(s in f for s in ("₹", "$", "[$", "Rs", "€", "£")) or "accounting" in f.lower(),
        "percent": lambda f: "%" in f,
        "decimal2": lambda f: "0.00" in f,
        "comma": lambda f: "#,##" in f,
        "date": lambda f: bool(re.search(r"[dmy]", f.lower())) and "general" not in f.lower(),
    }
    labels = {"currency": "Currency (₹)", "percent": "Percentage", "decimal2": "2 decimal places",
              "comma": "Comma style", "date": "Date"}

    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        if _all_empty(ws, rng):
            return False, f"{rng} is empty"
        bad = [f"{c.coordinate}: {c.number_format}" for c in cells(ws, rng) if c.value is not None and not tests[kind](c.number_format or "")]
        return (not bad), ("Format OK" if not bad else "Format: " + ", ".join(bad[:3]))
    return Check(f"{_sheet_label(sheet)}{rng} formatted as {labels[kind]}", marks, fn, expected=labels[kind],
                 how=f"Select > Home > Number group > choose {labels[kind]} (or Ctrl+1 > Number tab).", topic=T_FMT)


def col_width(sheet, col, width, marks=1.0, file=None, tol=1.0):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        idx = column_index_from_string(col)
        w = None
        for key, d in ws.column_dimensions.items():
            lo, hi = d.min or column_index_from_string(key), d.max or column_index_from_string(key)
            if lo <= idx <= hi and d.width:
                w = d.width
        if w is None:
            return False, "Column width not changed (default)"
        return abs(w - width) <= tol, f"Column {col} width {w:.1f}"
    return Check(f"Column {col} width about {width}", marks, fn, expected=f"Width {width}",
                 how=f"Right-click column header {col} > Column Width > {width} > OK.", topic="setting row height, column width")


def autofit(sheet, col, min_width, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        idx = column_index_from_string(col)
        w = None
        for key, d in ws.column_dimensions.items():
            lo, hi = d.min or column_index_from_string(key), d.max or column_index_from_string(key)
            if lo <= idx <= hi and d.width:
                w = d.width
        return (w or 0) >= min_width, f"Column {col} width {w or 'default'}"
    return Check(f"Column {col} widened to show its full text", marks, fn, expected=f"Width ≥ {min_width}",
                 how=f"Double-click the right edge of column header {col} (AutoFit).", topic="setting row height, column width")


def row_height(sheet, row, height, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        h = ws.row_dimensions[row].height
        if h is None:
            return False, "Row height not changed"
        return abs(h - height) <= 1, f"Row {row} height {h}"
    return Check(f"Row {row} height {height}", marks, fn, expected=str(height),
                 how=f"Right-click row number {row} > Row Height > {height} > OK.", topic="setting row height, column width")


def hidden_col(sheet, col, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        idx = column_index_from_string(col)
        for key, d in ws.column_dimensions.items():
            lo, hi = d.min or column_index_from_string(key), d.max or column_index_from_string(key)
            if lo <= idx <= hi and d.hidden:
                return True, f"Column {col} hidden"
        return False, f"Column {col} is visible"
    return Check(f"Column {col} hidden", marks, fn, expected="Hidden",
                 how=f"Right-click column header {col} > Hide.", topic="hiding/unhiding rows/columns")


def hidden_row(sheet, row, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        h = ws.row_dimensions[row].hidden
        return bool(h), f"Row {row} {'hidden' if h else 'visible'}"
    return Check(f"Row {row} hidden", marks, fn, expected="Hidden",
                 how=f"Right-click row number {row} > Hide.", topic="hiding/unhiding rows/columns")


def freeze(sheet, cell="A2", marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        return ws.freeze_panes == cell, f"Freeze panes: {ws.freeze_panes or 'none'}"
    return Check(f"Panes frozen at {cell}" + (" (top row)" if cell == "A2" else ""), marks, fn, expected=cell,
                 how="View > Freeze Panes > Freeze Top Row." if cell == "A2" else
                     f"Click {cell} > View > Freeze Panes > Freeze Panes.", topic=T_FMT)


# =================================================================== data tools

def sorted_by(sheet, col, first_row, last_row, desc=False, numeric=False, keep=None, marks=1.0, file=None):
    """keep = (key_col, other_col, mapping) to make sure whole rows moved together."""
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet, values=True)
        vals = [ws[f"{col}{r}"].value for r in range(first_row, last_row + 1)]
        if any(v is None for v in vals):
            return False, "Some rows are empty in the sort column"
        key = (lambda v: float(v)) if numeric else (lambda v: str(v).lower())
        try:
            ok = [key(v) for v in vals] == sorted((key(v) for v in vals), reverse=desc)
        except ValueError:
            return False, "Column has mixed text and numbers"
        if not ok:
            return False, "Not in order: " + short(", ".join(map(str, vals)), 60)
        if keep:
            kc, oc, mapping = keep
            for r in range(first_row, last_row + 1):
                k, o = ws[f"{kc}{r}"].value, ws[f"{oc}{r}"].value
                if k in mapping and _num(mapping[k]) != _num(o) and str(mapping[k]) != str(o):
                    return False, f"Sorted, but rows got mixed up (row {r}: {k} now shows {o}). Select the whole table before sorting."
        return True, "Sorted correctly"
    return Check(f"Data sorted by column {col} ({'largest/Z first' if desc else 'A→Z / smallest first'})", marks, fn,
                 expected="Rows in the right order, each row kept together",
                 how=f"Click any cell in the data > Data > Sort > Sort by column {col}, Order "
                     f"{'Largest to Smallest / Z to A' if desc else 'A to Z / Smallest to Largest'}, 'My data has headers' ticked > OK.",
                 topic="Sorting")


def autofilter(sheet, marks=1.0, file=None, applied=True):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        if not ws.auto_filter.ref:
            return False, "No filter buttons on the data"
        if applied:
            if not ws.auto_filter.filterColumn:
                return False, "Filter buttons added but no filter applied"
        return True, f"Filter on {ws.auto_filter.ref}" + (" with criteria" if applied else "")
    return Check("AutoFilter " + ("applied" if applied else "turned on"), marks, fn,
                 expected="Filter arrows on headers" + ("; only matching rows shown" if applied else ""),
                 how="Click in the data > Data > Filter (Ctrl+Shift+L) > click the arrow in the column > choose the value(s) > OK.",
                 topic="Filtering")


def validation(sheet, ref, kind="list", marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        for dv in ws.data_validations.dataValidation:
            if any(ref in str(r) or _in_range(ref, str(r)) for r in str(dv.sqref).split()):
                if dv.type == kind or kind == "any":
                    return True, f"Validation type '{dv.type}' on {dv.sqref}"
                return False, f"Validation type is '{dv.type}', not '{kind}'"
        x = b.sheet_xml(sheet)
        if "x14:dataValidation" in x:
            return True, "Validation found (list from another sheet)"
        return False, f"No data validation on {ref}"
    return Check(f"Data validation ({kind}) on {ref}", marks, fn, expected=f"{kind} validation",
                 how="Select the cells > Data > Data Validation > Allow: " +
                     ("List > Source: type the items separated by commas" if kind == "list" else "Whole number > set Minimum/Maximum") + " > OK.",
                 topic="Validation")


def _in_range(ref, rng):
    if ":" not in rng:
        return ref.upper() == rng.upper()
    try:
        a, b, c, d = range_boundaries(rng)
        col = column_index_from_string(re.match(r"[A-Z]+", ref.upper()).group(0))
        row = int(re.search(r"\d+", ref).group(0))
        return a <= col <= c and b <= row <= d
    except Exception:
        return False


def cond_format(sheet, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        n = sum(len(cf.rules) for cf in ws.conditional_formatting)
        if n == 0 and "conditionalFormatting" in b.sheet_xml(sheet):
            n = 1
        return n > 0, f"{n} conditional formatting rule(s)"
    return Check("Conditional formatting applied", marks, fn, expected="A highlight rule on the data",
                 how="Select the cells > Home > Conditional Formatting > Highlight Cells Rules > Greater Than > type value > OK.",
                 topic=T_FMT)


def subtotal(marks=2.0, file=None, sheet=None, min_groups=2):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        n = 0
        grand = False
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    if c.value.upper().startswith("=SUBTOTAL("):
                        n += 1
                    if "grand total" in c.value.lower():
                        grand = True
        ok = n >= min_groups and grand
        return ok, f"{n} SUBTOTAL formula(s); Grand Total row: {'yes' if grand else 'no'}"
    return Check("Subtotals created by group", marks, fn, expected="Total rows per group and a Grand Total",
                 how="First sort by the group column. Then Data > Subtotal > At each change in: the group column; "
                     "Use function: Sum; Add subtotal to: the amount column > OK.", topic="consolidation and subtotal")


def chart(kind=None, title=None, marks=1.0, file=None, what=None):
    tags = {"column": ["c:barChart", "c:bar3DChart"], "bar": ["c:barChart", "c:bar3DChart"], "pie": ["c:pieChart", "c:pie3DChart", "c:doughnutChart"],
            "line": ["c:lineChart", "c:line3DChart"]}

    def fn(ctx):
        b = ctx.excel(file)
        charts = [n for n in b.zip.namelist() if re.match(r"xl/charts/chart\d+\.xml", n)]
        if not charts:
            if any(n.startswith("xl/charts/chartEx") for n in b.zip.namelist()):
                return False, "Only a new-style chart (e.g. Histogram/Waterfall) found; use a normal Column/Bar/Pie/Line chart"
            return False, "No chart found"
        msgs = []
        for n in charts:
            x = b.zip.read(n).decode("utf8", "ignore")
            types = re.findall(r"<(c:\w+Chart)>", x) + re.findall(r"<(c:\w+Chart) ", x)
            ctype = types[0] if types else "?"
            if kind:
                if not any(t in x for t in tags[kind]):
                    msgs.append(f"chart is {ctype.replace('c:', '')}")
                    continue
                if kind in ("column", "bar"):
                    m = re.search(r'<c:barDir val="(\w+)"', x)
                    d = m.group(1) if m else "col"
                    if (kind == "column") != (d == "col"):
                        msgs.append(f"chart is a {'bar' if d == 'bar' else 'column'} chart")
                        continue
            if title:
                t = " ".join(re.findall(r"<a:t>([^<]*)</a:t>", x))
                if norm(title) not in norm(t):
                    msgs.append(f"title is '{short(t, 30) or 'missing'}'")
                    continue
            return True, f"{(kind or 'A').capitalize()} chart found" + (f" with title '{title}'" if title else "")
        return False, "; ".join(msgs)
    return Check(what or f"{(kind or 'A').capitalize()} chart" + (f" titled '{title}'" if title else ""), marks, fn,
                 expected=f"{kind or ''} chart" + (f" with title {title}" if title else ""),
                 how=f"Select the data (names and values) > Insert > {'Column' if kind == 'column' else (kind or 'Column').capitalize()} > "
                     "first 2-D style. Click the chart title and type the new title.", topic="creating charts")


def chart_data_labels(marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        x = b.xml("xl/charts/chart")
        ok = '<c:showVal val="1"' in x or '<c:showPercent val="1"' in x
        return ok, "Data labels shown" if ok else "No data labels"
    return Check("Chart shows data labels", marks, fn, expected="Values on the chart",
                 how="Click the chart > Chart Elements (+) > Data Labels. (2010: Layout > Data Labels.)", topic="creating charts")


# =================================================================== sheets

def sheet_exists(name, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ok = name.lower() in [s.lower() for s in b.wb.sheetnames]
        return ok, "Sheets: " + ", ".join(b.wb.sheetnames)
    return Check(f"Sheet named '{name}'", marks, fn, expected=name,
                 how="Double-click the sheet tab (or right-click > Rename) > type the name > Enter. "
                     "New sheet: click the + next to the tabs (Shift+F11).", topic=T_SHEETS)


def sheet_absent(name, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ok = name.lower() not in [s.lower() for s in b.wb.sheetnames]
        return ok, "Sheets: " + ", ".join(b.wb.sheetnames)
    return Check(f"Sheet '{name}' deleted", marks, fn, expected=f"No sheet called {name}",
                 how=f"Right-click the '{name}' tab > Delete.", topic=T_SHEETS)


def sheet_first(name, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        return b.wb.sheetnames[0].lower() == name.lower(), "Order: " + ", ".join(b.wb.sheetnames)
    return Check(f"Sheet '{name}' moved to first position", marks, fn, expected=f"{name} first",
                 how="Drag the tab to the far left (or right-click > Move or Copy > (move to beginning)).", topic=T_SHEETS)


def tab_color(sheet, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        c = ws.sheet_properties.tabColor
        return c is not None, "Tab colour set" if c is not None else "No tab colour"
    return Check(f"Tab colour on '{sheet}'", marks, fn, expected="Coloured tab",
                 how="Right-click the sheet tab > Tab Color > pick a colour.", topic=T_SHEETS)


def comment(sheet, ref, marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        ws = b.ws(sheet)
        ok = ws[ref].comment is not None
        if not ok and any("threadedComment" in n for n in b.zip.namelist()):
            ok = True
        return ok, "Comment found" if ok else f"No comment on {ref}"
    return Check(f"Comment/Note on {ref}", marks, fn, expected="A comment",
                 how=f"Right-click {ref} > Insert Comment (New Note in 365) > type > click outside.", topic="Inserting")


def hyperlink(sheet, ref, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        h = ws[ref].hyperlink
        f = ws[ref].value
        ok = h is not None or (isinstance(f, str) and f.upper().startswith("=HYPERLINK("))
        return ok, "Hyperlink found" if ok else "No hyperlink"
    return Check(f"Hyperlink in {ref}", marks, fn, expected="Clickable link",
                 how=f"Click {ref} > Ctrl+K > type the address > OK.", topic="Inserting")


def picture(marks=1.0, file=None):
    def fn(ctx):
        b = ctx.excel(file)
        n = len([x for x in b.zip.namelist() if x.startswith("xl/media/")])
        return n > 0, f"{n} picture(s)"
    return Check("Picture/clip art inserted", marks, fn, expected="A picture on the sheet",
                 how="Insert > Pictures > choose an image from the test folder > Insert.", topic="clipart, pictures")


def symbol(sheet, ref, chars="₹©®", marks=1.0, file=None):
    def fn(ctx):
        v = str(ctx.excel(file).ws(sheet)[ref].value or "")
        ok = any(c in v for c in chars)
        return ok, f"{ref} = {v}"
    return Check(f"Symbol inserted in {ref}", marks, fn, expected=" or ".join(chars),
                 how="Insert > Symbol > pick the symbol > Insert.", topic="symbol")


def protected(sheet, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        return bool(ws.protection.sheet), "Sheet protected" if ws.protection.sheet else "Sheet not protected"
    return Check(f"Sheet '{sheet}' protected", marks, fn, expected="Protected sheet",
                 how="Review > Protect Sheet > (password optional) > OK.", topic=T_SHEETS)


# =================================================================== printing

def landscape(sheet, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        return ws.page_setup.orientation == "landscape", f"Orientation: {ws.page_setup.orientation or 'portrait'}"
    return Check("Page orientation Landscape (print setup)", marks, fn, expected="Landscape",
                 how="Page Layout > Orientation > Landscape.", topic="Printing")


def fit_one_page(sheet, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        fit = ws.sheet_properties.pageSetUpPr is not None and ws.sheet_properties.pageSetUpPr.fitToPage
        return bool(fit), "Fit to page on" if fit else "Not set to fit on one page"
    return Check("Set to print on one page", marks, fn, expected="Fit Sheet on One Page",
                 how="File > Print > Scaling: 'Fit Sheet on One Page' (or Page Layout > Width 1 page, Height 1 page). Then save.",
                 topic="Printing")


def print_titles(sheet, rows="1:1", marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        t = str(ws.print_title_rows or "")
        ok = t.replace("$", "") == rows
        return ok, f"Print titles: {t or 'none'}"
    return Check(f"Row {rows.split(':')[0]} repeats on every printed page", marks, fn, expected=f"Rows to repeat at top: ${rows.replace(':', ':$')}",
                 how="Page Layout > Print Titles > Rows to repeat at top: click row 1 > OK.", topic="Printing")


def print_area(sheet, rng, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        a = str(ws.print_area or "")
        ok = rng.replace("$", "").upper() in a.replace("$", "").upper()
        return ok, f"Print area: {a or 'not set'}"
    return Check(f"Print area set to {rng}", marks, fn, expected=rng,
                 how=f"Select {rng} > Page Layout > Print Area > Set Print Area.", topic="Printing")


def sheet_header(sheet, text, marks=1.0, file=None, footer=False):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        hf = ws.oddFooter if footer else ws.oddHeader
        t = " ".join(filter(None, [hf.left.text, hf.center.text, hf.right.text]))
        ok = norm(text) in norm(t) if text else bool(t)
        return ok, f"{'Footer' if footer else 'Header'}: '{t or 'empty'}'"
    return Check(f"Print {'footer' if footer else 'header'} contains '{text}'", marks, fn, expected=text,
                 how=f"Insert > Header & Footer > click in the {'footer' if footer else 'header'} > type > click a cell. "
                     "(Or Page Setup > Header/Footer tab > Custom.)", topic="Printing")


def page_break(sheet, marks=1.0, file=None):
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet)
        n = len(ws.row_breaks.brk) + len(ws.col_breaks.brk)
        return n > 0, f"{n} manual page break(s)"
    return Check("Manual page break inserted", marks, fn, expected="A page break",
                 how="Click the row where the new page should start > Page Layout > Breaks > Insert Page Break.",
                 topic="manual breaks")


def rowwise(sheet, col, first, last, func, compute, marks=1.0, file=None, what=None, how=None, tol=0.01,
            topic=T_FORMULA):
    """Each row's formula in `col` must use `func` and give compute(row) where row maps column letters to values.
    Works even after the rows have been sorted."""
    def fn(ctx):
        b = ctx.excel(file)
        ws, wv = b.ws(sheet), b.ws(sheet, values=True)
        bad = []
        for r in range(first, last + 1):
            f = ws[f"{col}{r}"].value
            if f is None:
                bad.append(f"{col}{r} is empty")
                continue
            if not (isinstance(f, str) and f.startswith("=")):
                bad.append(f"{col}{r} has a typed value ({f}), not a formula")
                continue
            if func and not re.search(r"\b" + func + r"\s*\(" if func.isalpha() else func, f, re.I):
                bad.append(f"{col}{r} = {f}")
                continue
            row = {get_column_letter(c): wv.cell(r, c).value for c in range(1, wv.max_column + 1)}
            try:
                want = compute(row)
            except Exception:
                continue
            got = wv[f"{col}{r}"].value
            if got is None:
                continue                      # file saved without calculated values; formula itself is fine
            if isinstance(want, (int, float)):
                n = _num(got)
                if n is None or abs(n - want) > tol:
                    bad.append(f"{col}{r} gives {got} (should be {round(want, 2):g})")
            elif norm(str(got)) != norm(str(want)):
                bad.append(f"{col}{r} gives {got} (should be {want})")
        n = last - first + 1
        return partial(n - len(bad), n), (f"All {n} rows correct" if not bad else f"{len(bad)} of {n} rows wrong: " + "; ".join(bad[:3]))
    return Check(what or f"{_sheet_label(sheet)}{col}{first}:{col}{last} use {func}", marks, fn,
                 expected=(f"A {func} formula" if func.isalpha() else "A formula") + f" in every row of column {col}",
                 how=how or f"Type the formula in {col}{first}, press Enter, then drag the fill handle down to row {last}.",
                 topic=topic)


def labelled_values(sheet, label_col, value_col, mapping, marks=1.0, file=None, what=None, how="", topic=T_FORMULA,
                    tol=0.01):
    """For each label in mapping, some row must have that label in label_col and the value in value_col."""
    def fn(ctx):
        ws = ctx.excel(file).ws(sheet, values=True)
        found = {}
        for r in range(1, ws.max_row + 1):
            lab = ws[f"{label_col}{r}"].value
            if lab is not None:
                found.setdefault(norm(str(lab)), ws[f"{value_col}{r}"].value)
        bad = []
        for lab, want in mapping.items():
            got = found.get(norm(lab))
            n = _num(got)
            if n is None or abs(n - want) > tol:
                bad.append(f"{lab}: {got if got is not None else 'missing'} (should be {want:g})")
        return partial(len(mapping) - len(bad), len(mapping)), ("All values correct" if not bad else "; ".join(bad[:4]))
    return Check(what or f"Correct totals on {sheet.split('|')[0]}", marks, fn,
                 expected="; ".join(f"{k} = {v:g}" for k, v in list(mapping.items())[:5]), how=how, topic=topic)
